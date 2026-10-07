import os
import sys
import re
import time
import requests
from pathlib import Path
from dotenv import load_dotenv

# Ensure project root is in sys.path
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from google import genai
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

FONT_NAME = "Kantumruy Pro"

def get_vertex_client():
    """Fetches credentials from Bifrost/environment and initializes Vertex AI Client."""
    load_dotenv()
    load_dotenv(Path(__file__).resolve().parent.parent / ".env")
    
    bifrost_url = os.getenv("BIFROST_URL")
    client_id = os.getenv("BIFROST_CLIENT_ID")
    webhook_secret = os.getenv("BIFROST_WEBHOOK_SECRET")
    
    sa_json = None
    project = "khmer-ocr-496606"
    location = "asia-southeast1"
    
    if bifrost_url and client_id and webhook_secret:
        try:
            res = requests.get(
                f"{bifrost_url.rstrip('/')}/api/v1/config",
                headers={"X-Client-ID": client_id, "X-Webhook-Secret": webhook_secret},
                timeout=10
            )
            if res.status_code == 200:
                data = res.json().get("data", {}).get("api_keys", {})
                sa_json = data.get("GOOGLE_APPLICATION_CREDENTIALS_JSON")
                project = data.get("VERTEX_AI_PROJECT", project)
                location = data.get("VERTEX_AI_LOCATION", location)
        except Exception as e:
            print(f"Warning: Could not fetch secrets from Bifrost: {e}")

    if sa_json:
        sa_path = "/tmp/bifrost_vertex_sa.json"
        with open(sa_path, "w", encoding="utf-8") as f:
            f.write(sa_json)
        os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = sa_path

    return genai.Client(vertexai=True, project=project, location=location)

def set_font_properties(run, font_name=FONT_NAME, size_pt=11, bold=False, italic=False, color_rgb=None):
    """Sets Latin and Complex Script (Khmer) font properties on a docx Run object."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:cs'), font_name)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)

def fix_khmer_spacing(text: str) -> str:
    """Removes unnatural spaces between Khmer characters (Khmer script does not use spaces between words)."""
    prev = None
    curr = text
    while prev != curr:
        prev = curr
        curr = re.sub(r'([\u1780-\u17FF\u17D2])\s+([\u1780-\u17FF])', r'\1\2', curr)
    return curr

def split_text_into_blocks(text: str, max_chars: int = 2500) -> list[str]:
    """Splits text cleanly into sub-blocks of max_chars length for rapid LLM processing."""
    blocks = []
    i = 0
    n = len(text)
    while i < n:
        end = min(i + max_chars, n)
        if end < n:
            space_idx = text.rfind(" ", i + max_chars - 400, end)
            if space_idx > i:
                end = space_idx
        blocks.append(text[i:end].strip())
        i = end
    return [b for b in blocks if b]

def clean_chronological_ideas(client, chunk_raw_text: str, model_name: str = "gemini-2.5-flash") -> str:
    """Cleans transcript chunk chronologically idea-by-idea into simplified, natural Khmer prose."""
    prompt = (
        "You are an expert Khmer language editor.\n"
        "Below is a raw verbatim transcript of spoken dialogue from a recording.\n\n"
        "Your task is to simplify and clean up what was spoken IDEA BY IDEA in exact CHRONOLOGICAL ORDER:\n"
        "1. STRICT CHRONOLOGICAL SEQUENCE: Follow the spoken order strictly as said in the audio. Do NOT re-order topics or write a report summary.\n"
        "2. SIMPLIFY IDEA BY IDEA: Take whatever idea was expressed and rewrite it into simple, clear, elegant, natural Khmer sentences.\n"
        "3. REMOVE FILLER & STUTTERS: Strip out filler words ('អឺ', 'អា', 'ហ្នឹង', 'អញ្ចឹងទៅ', 'បាទ/ចាស', 'គេហៅថា'), speech stutters, false starts, and repeated words.\n"
        "4. KHMER SPACING RULE: In standard Khmer writing, write continuous Khmer text WITHOUT spaces between words in a phrase. Only put spaces around English terms or between major sentences/clauses.\n"
        "5. PRESERVE TECHNICAL TERMS: Keep all technical and business English terms (API, CamDX, Single Portal, OBR, ESB, UI/UX, Q&A, SBI, Startup, etc.) in clean English text.\n\n"
        "Output ONLY the cleaned chronological text."
    )
    
    blocks = split_text_into_blocks(chunk_raw_text, max_chars=2500)
    cleaned_blocks = []
    
    for idx, b in enumerate(blocks, 1):
        if not b.strip():
            continue
        print(f"    Sub-block {idx}/{len(blocks)} ({len(b)} chars)...")
        max_retries = 3
        cleaned_b = ""
        for attempt in range(1, max_retries + 1):
            try:
                res = client.models.generate_content(
                    model=model_name,
                    contents=[prompt, b]
                )
                if res.text:
                    cleaned_b = fix_khmer_spacing(res.text.strip())
                    break
            except Exception as e:
                print(f"    Attempt {attempt} sub-block cleaning error: {e}")
                time.sleep(2)
        if not cleaned_b:
            cleaned_b = fix_khmer_spacing(b)
        cleaned_blocks.append(cleaned_b)

    return "\n\n".join(cleaned_blocks)

def export_chronological_cleaned_docx(txt_files: list[str], output_docx: str):
    """Cleans transcripts idea-by-idea chronologically and exports a Kantumruy Pro DOCX file."""
    client = get_vertex_client()
    doc = docx.Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Main Document Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(8)
    t_run = title_p.add_run("ប្រតិចារិកសម្រួល និងសម្អាតតាមលំដាប់លំដោយ - Au Veng Street")
    set_font_properties(t_run, font_name=FONT_NAME, size_pt=20, bold=True, color_rgb=RGBColor(31, 73, 125))

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(24)
    s_run = sub_p.add_run("Chronological Cleaned Transcript (Idea-by-Idea in Kantumruy Pro)")
    set_font_properties(s_run, font_name=FONT_NAME, size_pt=12, italic=True, color_rgb=RGBColor(89, 89, 89))

    for txt_file in txt_files:
        path = Path(txt_file)
        if not path.exists():
            continue

        print(f"\nProcessing & Cleaning Chronologically: {path.name}...")
        raw_content = path.read_text(encoding="utf-8")
        
        chunks = raw_content.split("--- Chunk ")

        file_h = doc.add_paragraph()
        file_h.paragraph_format.space_before = Pt(18)
        file_h.paragraph_format.space_after = Pt(10)
        fh_run = file_h.add_run(f"📋 {path.stem.replace('_transcription', '')}")
        set_font_properties(fh_run, font_name=FONT_NAME, size_pt=16, bold=True, color_rgb=RGBColor(54, 95, 145))

        for chunk in chunks:
            if not chunk.strip():
                continue

            chunk_parts = chunk.split("\n\n", 1)
            chunk_num = chunk_parts[0].split(" ---")[0] if " ---" in chunk_parts[0] else ""
            chunk_text = chunk_parts[1] if len(chunk_parts) > 1 else chunk_parts[0]

            print(f"  Cleaning Chunk {chunk_num} chronologically (length: {len(chunk_text)} chars)...")
            cleaned_text = clean_chronological_ideas(client, chunk_text)

            if chunk_num:
                cp = doc.add_paragraph()
                cp.paragraph_format.space_before = Pt(12)
                cp.paragraph_format.space_after = Pt(6)
                c_run = cp.add_run(f"--- ផ្នែកទី {chunk_num} (Part {chunk_num}) ---")
                set_font_properties(c_run, font_name=FONT_NAME, size_pt=12, bold=True, color_rgb=RGBColor(89, 89, 89))

            lines = cleaned_text.splitlines()
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue

                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_after = Pt(6)

                if line_str.startswith("# "):
                    p.paragraph_format.space_before = Pt(12)
                    run = p.add_run(line_str.lstrip("# ").strip())
                    set_font_properties(run, font_name=FONT_NAME, size_pt=14, bold=True, color_rgb=RGBColor(31, 73, 125))
                elif line_str.startswith("## "):
                    p.paragraph_format.space_before = Pt(10)
                    run = p.add_run(line_str.lstrip("## ").strip())
                    set_font_properties(run, font_name=FONT_NAME, size_pt=13, bold=True, color_rgb=RGBColor(54, 95, 145))
                elif line_str.startswith("- ") or line_str.startswith("* "):
                    p.paragraph_format.left_indent = Inches(0.25)
                    run = p.add_run("• " + line_str[2:].strip())
                    set_font_properties(run, font_name=FONT_NAME, size_pt=11)
                else:
                    run = p.add_run(line_str)
                    set_font_properties(run, font_name=FONT_NAME, size_pt=11)

    out_path = Path(output_docx).resolve()
    doc.save(str(out_path))
    print(f"\n🎉 CHRONOLOGICAL CLEANED DOCX SAVED TO: {out_path}")
    return str(out_path)

if __name__ == "__main__":
    sys.stdout.reconfigure(line_buffering=True)
    txt_inputs = [
        "/Users/nicksng/code/random/Au Veng Street 3_transcription.txt",
        "/Users/nicksng/code/random/Au Veng Street 4_transcription.txt"
    ]
    output_docx = "/Users/nicksng/code/random/Au_Veng_Street_Chronological_Cleaned_Transcript.docx"
    export_chronological_cleaned_docx(txt_inputs, output_docx)
