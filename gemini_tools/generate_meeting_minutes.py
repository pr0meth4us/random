import os
import sys
import re
import time
import requests
from pathlib import Path
from dotenv import load_dotenv

# Ensure project root is in sys.path
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from google import genai
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
import pptx

FONT_NAME = "Kantumruy Pro"

def get_vertex_client():
    """Fetches credentials from Bifrost/environment and initializes Vertex AI Client."""
    load_dotenv()
    load_dotenv(Path(__file__).resolve().parent.parent / ".env")
    
    bifrost_url = os.getenv("BIFROST_URL")
    client_id = os.getenv("BIFROST_CLIENT_ID")
    webhook_secret = os.getenv("BIFROST_WEBHOOK_SECRET")
    
    sa_json = None
    project = "khmer-ocr-496606"
    location = "asia-southeast1"
    
    if bifrost_url and client_id and webhook_secret:
        try:
            res = requests.get(
                f"{bifrost_url.rstrip('/')}/api/v1/config",
                headers={"X-Client-ID": client_id, "X-Webhook-Secret": webhook_secret},
                timeout=10
            )
            if res.status_code == 200:
                data = res.json().get("data", {}).get("api_keys", {})
                sa_json = data.get("GOOGLE_APPLICATION_CREDENTIALS_JSON")
                project = data.get("VERTEX_AI_PROJECT", project)
                location = data.get("VERTEX_AI_LOCATION", location)
        except Exception as e:
            print(f"Warning: Could not fetch secrets from Bifrost: {e}")

    if sa_json:
        sa_path = "/tmp/bifrost_vertex_sa.json"
        with open(sa_path, "w", encoding="utf-8") as f:
            f.write(sa_json)
        os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = sa_path

    return genai.Client(vertexai=True, project=project, location=location)

def set_font_properties(run, font_name=FONT_NAME, size_pt=11, bold=False, italic=False, color_rgb=None):
    """Sets Latin and Complex Script (Khmer) font properties on a docx Run object."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:cs'), font_name)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)

def fix_khmer_spacing(text: str) -> str:
    """Removes unnatural spaces between Khmer characters (Khmer script does not use spaces between words)."""
    prev = None
    curr = text
    while prev != curr:
        prev = curr
        curr = re.sub(r'([\u1780-\u17FF\u17D2])\s+([\u1780-\u17FF])', r'\1\2', curr)
    return curr

def extract_pptx_slides_content(pptx_path: str) -> str:
    """Extracts text content slide by slide from PPTX presentation file."""
    prs = pptx.Presentation(pptx_path)
    output = []
    for idx, slide in enumerate(prs.slides, 1):
        output.append(f"--- SLIDE {idx} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        output.append(f"- {t}")
    return "\n".join(output)

def generate_meeting_minutes_from_flow(client, pptx_content: str, full_transcript: str, model_name: str = "gemini-2.5-flash") -> str:
    """Generates official Meeting Minutes (កំណត់ហេតុប្រជុំ) merging ALL spoken transcript content into the presentation slide flow."""
    prompt = (
        "You are an elite secretary general and official rapporteur in Khmer & English for Cambodian digital government committees.\n\n"
        "Below are TWO SOURCES OF INFORMATION:\n\n"
        "1. PRESENTATION SLIDES & AGENDA FLOW (from PPTX file):\n" + pptx_content + "\n\n"
        "2. COMPLETE AUDIO MEETING TRANSCRIPTIONS (Au Veng Street 3 and 4):\n" + full_transcript + "\n\n"
        "CRITICAL INSTRUCTION: MERGE ALL SPOKEN TRANSCRIPT CONTENT DIRECTLY INTO THE PRESENTATION SLIDE AGENDA FLOW to produce a comprehensive, exhaustive MEETING MINUTES DOCUMENT (កំណត់ហេតុប្រជុំតម្រង់ទិស និងយន្តការការងារក្រុមការងារ) in Khmer.\n\n"
        "REQUIREMENTS:\n"
        "- Under each section of the Presentation Slide Agenda, include ALL specific discussion points, speaker explanations, technical decisions, questions asked, and answers provided during the meeting.\n"
        "- Do NOT omit any spoken details (e.g. 733 to 1,000+ services by 2026, 50% fee reduction, Bakong bank integration, NRMIS auto-invoice, verify.gov.kh, weekly Friday 9:00 AM reporting to PM, containerized shell templates, tutorial videos, soft vs official launch).\n\n"
        "STRUCTURE TO FOLLOW (Mapped to Presentation Flow):\n\n"
        "# កំណត់ហេតុប្រជុំតម្រង់ទិស និងយន្តការការងាររបស់ក្រុមការងារអន្តរក្រសួង\n"
        "*(Merged Presentation Agenda & Complete Spoken Discussion Minutes)*\n\n"
        "## ១. របៀបវារៈ និងគោលបំណងនៃកិច្ចប្រជុំ (Meeting Agenda & Objectives - Slide 1)\n"
        "- Detailed explanation of the meeting purpose, onboarding of taskforce members, and context.\n\n"
        "## ២. រំហូរការងារជាមួយក្រសួង-ស្ថាប័ន (Workflow with Line Ministries - Slide 2)\n"
        "### ២.១. ការកំណត់ និងចាត់ថ្នាក់សេវាពាក់ព័ន្ធ (Service Identification & Categorization)\n"
        "- ករណីស្ថាប័នមានប្រព័ន្ធ (Has Existing System): ធ្វើសន្ធានកម្មមកថ្នាល CamDX.\n"
        "- ករណីស្ថាប័នគ្មានប្រព័ន្ធ (No System): សម្របសម្រួលអភិវឌ្ឍប្រព័ន្ធមកថ្នាល CamDX ដោយប្រើ ESB Generic Template (សំបកប្រព័ន្ធ).\n"
        "### ២.២. កិច្ចការលម្អិតរបស់ក្រុមការងារសម្របសម្រួល (Coordination Taskforce Detailed Roles)\n"
        "- ពិនិត្យទម្រង់ពាក្យស្នើសុំ ការបង្រួម/ផ្គួប Data field និងឯកសារភ្ជាប់ (ទិន្នន័យដែលមានស្រាប់ក្នុង CamDX, ទិន្នន័យមិនមានតម្លៃបន្ថែមខ្ពស់).\n"
        "- ពិនិត្យនីតិវិធី និងរំហូរការងារ: ការកាត់បន្ថយរយៈពេល និងកម្រៃសេវា (គោលការណ៍កាត់បន្ថយ ៥០%).\n"
        "- ការប្រើប្រាស់ UX/UI របស់ Generic Template.\n"
        "### ២.៣. កិច្ចការលម្អិតរបស់ក្រុមការងារអភិវឌ្ឍ (Development Taskforce Detailed Roles)\n"
        "- ពិនិត្យទម្រង់ពាក្យស្នើសុំ និង Data Points អនុញ្ញាតឱ្យទាញទិន្នន័យ.\n"
        "- ការបើក API ជាមួយថ្នាល CamDX លើ ៣ មុខងារសំខាន់ៗ: (1) Application (ទិន្នន័យដុលពាក្យស្នើសុំ), (2) License (ទិន្នន័យសកម្ម: Total, Approved, Rejected, Returned, Pending, Revoked), (3) Statistics (ទិន្នន័យ Accumulative សម្រាប់រាយការណ៍).\n"
        "- ការធានានូវគោលការណ៍ DPIs ទាំង ៤.\n"
        "### ២.៤. ការពិនិត្យ សម្រេច និងការងារ PR (Review, Approval & PR Taskforce)\n"
        "- សម្រេចឯកភាពដោយក្រសួង/ស្ថាប័នសាម៉ី.\n"
        "- ក្រុមការងារ PR: សេចក្តីប្រកាសព័ត៌មាន, ផ្ទាំងផ្សព្វផ្សាយ, វីដេអូណែនាំសម្រាប់សាធារណជន.\n"
        "- យុទ្ធសាស្ត្រដាក់ឱ្យដំណើរការ: Soft Launch និង Official Launch (ផ្តល់ Spotlight ជូនក្រសួងសាម៉ី).\n\n"
        "## ៣. លម្អិតភារកិច្ចក្រុមការងារ និងការអនុវត្ត DPI (Detailed Tasks & DPI Implementation - Slide 3)\n"
        "### ៣.១. ភារកិច្ចលម្អិតក្រុមការងារសម្របសម្រួល\n"
        "- ការពិនិត្យពាក្យស្នើសុំ ឯកសារភ្ជាប់ និងលិខិតបទដ្ឋានគតិយុត្តិ (បើក, បន្តសុពលភាព, ធ្វើបច្ចុប្បន្នភាព, បិទសេវា).\n"
        "- ការគ្រប់គ្រងរំហូរការងារ (ទទួលពាក្យ, បញ្ជូនត្រឡប់, បញ្ជូនបន្ត, បដិសេធ, អនុម័ត, ជូនដំណឹងផុតកំណត់វិញ្ញាបនបត្រ).\n"
        "- ការបង្កើតក្រុមការងារបង្គោលជាមួយក្រសួង-ស្ថាប័ន.\n"
        "### ៣.២. ការអនុវត្តតាមគោលការណ៍ public infrastructure ឌីជីថល (DPI Principles)\n"
        "- ការតភ្ជាប់ជាមួយធនាគារសមាជិក Bakong (API integration, documents, MoU ជាមួយធនាគារ).\n"
        "- ការតភ្ជាប់ជាមួយប្រព័ន្ធ NRMIS (លេខកូដសេវា, Generate Invoice ស្វ័យប្រវត្តិ, Bank Reference Settlement Number).\n"
        "- ការស្នើសុំប្រើប្រាស់ verify.gov.kh ពីគណៈកម្មាធិការរដ្ឋាភិបាលឌីជីថល (DGC).\n\n"
        "## ៤. ណែនាំ ESB Guidelines និងការសាកល្បងរបៀបការងារថ្មី (ESB Guidelines & Work Methodology - Slide 4)\n"
        "- ការរៀបចំ Figma Prototype ឱ្យដូចប្រព័ន្ធពិតដើម្បីជជែក និងទទួលយោបល់កែលម្អពីក្រសួង-ស្ថាប័ន មុនពេលអភិវឌ្ឍប្រព័ន្ធ.\n"
        "- របៀបសាកល្បងការងារថ្មី លើការជជែកនីតិវិធី និងបច្ចេកទេស.\n\n"
        "## ៥. យន្តការរាយការណ៍ប្រចាំសប្ដាហ៍ និងការបែងចែកភារកិច្ច (Weekly Reporting & Task Delegation)\n"
        "- យន្តការរាយការណ៍រៀងរាល់ព្រឹកថ្ងៃសុក្រ ម៉ោង ៩:០០ ព្រឹក ជូនថ្នាក់ដឹកនាំ និងសម្ដេចធិបតី (ប្រើប្រាស់ Bullet Points Reporting Template ខ្លី ខ្លឹម Clean).\n"
        "- ការចាត់តាំងតំណាងក្រុមការងាររៀបចំរបាយការណ៍ និងជំហានអនុវត្តបន្ត.\n\n"
        "WRITING RULES:\n"
        "1. KHMER SPACING RULE: Write natural continuous Khmer text without word spaces inside phrases. Only put spaces around English terms or major sentence boundaries.\n"
        "2. PRESERVE TECHNICAL TERMS: Keep all technical and business terms (API, CamDX, Single Portal, OBR, ESB, UI/UX, DPI, NRMIS, Bakong, verify.gov.kh, Figma, MoU, PR, Q&A) in clean English text.\n\n"
        "Output ONLY the detailed Markdown minutes document."
    )

    max_retries = 3
    for attempt in range(1, max_retries + 1):
        try:
            res = client.models.generate_content(
                model=model_name,
                contents=[prompt]
            )
            if res.text:
                return fix_khmer_spacing(res.text.strip())
        except Exception as e:
            print(f"  Attempt {attempt} meeting minutes generation error: {e}")
            time.sleep(3)
            
    return fix_khmer_spacing(pptx_content)

def export_meeting_minutes_docx(pptx_path: str, txt_files: list[str], output_docx: str):
    """Generates official Meeting Minutes in Kantumruy Pro DOCX matching PPTX flow."""
    client = get_vertex_client()
    doc = docx.Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    print("\n==================================================")
    print(" Extracting PPTX Presentation Flow...")
    print(f" Source PPTX: {Path(pptx_path).name}")
    print("==================================================\n")

    pptx_text = extract_pptx_slides_content(pptx_path)

    combined_transcripts = []
    for txt_file in txt_files:
        path = Path(txt_file)
        if path.exists():
            combined_transcripts.append(f"\n--- Transcript: {path.name} ---\n" + path.read_text(encoding="utf-8"))
    
    full_transcript = "\n".join(combined_transcripts)

    print("==================================================")
    print(" Merging Presentation Flow with Transcript Content...")
    print(" Engine: Gemini 2.5 Flash")
    print(" Font Family: Kantumruy Pro")
    print("==================================================\n")

    minutes_md = generate_meeting_minutes_from_flow(client, pptx_text, full_transcript)

    # Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(8)
    t_run = title_p.add_run("កំណត់ហេតុប្រជុំតម្រង់ទិស និងយន្តការការងាររបស់ក្រុមការងារអន្តរក្រសួង")
    set_font_properties(t_run, font_name=FONT_NAME, size_pt=20, bold=True, color_rgb=RGBColor(31, 73, 125))

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(24)
    s_run = sub_p.add_run("Merged Presentation Flow & Transcript Meeting Minutes (Kantumruy Pro)")
    set_font_properties(s_run, font_name=FONT_NAME, size_pt=12, italic=True, color_rgb=RGBColor(89, 89, 89))

    lines = minutes_md.splitlines()
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(6)

        if line_str.startswith("# "):
            continue
        elif line_str.startswith("## "):
            p.paragraph_format.space_before = Pt(16)
            run = p.add_run(line_str.lstrip("## ").strip())
            set_font_properties(run, font_name=FONT_NAME, size_pt=14, bold=True, color_rgb=RGBColor(31, 73, 125))
        elif line_str.startswith("### "):
            p.paragraph_format.space_before = Pt(10)
            run = p.add_run(line_str.lstrip("### ").strip())
            set_font_properties(run, font_name=FONT_NAME, size_pt=12, bold=True, color_rgb=RGBColor(54, 95, 145))
        elif line_str.startswith("- ") or line_str.startswith("* "):
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run("• " + line_str[2:].strip())
            set_font_properties(run, font_name=FONT_NAME, size_pt=11)
        else:
            run = p.add_run(line_str)
            set_font_properties(run, font_name=FONT_NAME, size_pt=11)

    out_path = Path(output_docx).resolve()
    doc.save(str(out_path))
    print(f"\n🎉 MERGED MEETING MINUTES DOCX SAVED TO: {out_path}")
    return str(out_path)

if __name__ == "__main__":
    sys.stdout.reconfigure(line_buffering=True)
    pptx_file = "/Users/nicksng/code/random/Orientation Taskforce resources .pptx"
    txt_inputs = [
        "/Users/nicksng/code/random/Au Veng Street 3_transcription.txt",
        "/Users/nicksng/code/random/Au Veng Street 4_transcription.txt"
    ]
    output_docx = "/Users/nicksng/code/random/Orientation_Taskforce_Meeting_Minutes.docx"
    export_meeting_minutes_docx(pptx_file, txt_inputs, output_docx)
