"""Delete text in a .docx as Word *tracked changes* (w:del), leaving everything else untouched.

    python docx_track_delete.py IN.docx OUT.docx --pattern REGEX [--author NAME]
    python docx_track_delete.py --selftest

Library use: track_delete(src, dst, find_spans, author) where find_spans(paragraph_text) returns
[(start, end), ...] to delete in that paragraph's *current* text (accepted view).

Safety rules:
- Only plain runs (direct children of w:p holding just w:t) are edited. Text inside existing
  insertions, hyperlinks, fields, or runs with tabs/breaks/drawings acts as a barrier: a span
  touching it is skipped and reported, never half-applied.
- Text inside existing w:del is ignored (it is not part of the current text).
- Revision ids continue after the highest existing w:id.
"""
import argparse, copy, re, sys, zipfile
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn

BARRIER = "￿"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _plain(r):
    kids = [k for k in r if k.tag != qn("w:rPr")]
    return bool(kids) and all(k.tag == qn("w:t") for k in kids)


def _segments(p):
    """[(run_or_None, text)] in document order; None marks barrier text."""
    out = []
    for child in p:
        if child.tag == qn("w:r"):
            txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
            out.append((child if _plain(child) else None, txt if _plain(child) else BARRIER * max(1, len(txt))))
        elif child.tag == qn("w:del"):
            continue
        elif child.tag in (qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"),
                           qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:proofErr")):
            continue
        else:
            txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if txt:
                out.append((None, BARRIER * len(txt)))
    return out


def _set_text(r, text):
    for t in list(r.iter(qn("w:t"))):
        r.remove(t)
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set(XML_SPACE, "preserve")
    r.append(t)


def _split(r, offset):
    """Split plain run at offset; returns right-hand run inserted after r."""
    text = "".join(t.text or "" for t in r.iter(qn("w:t")))
    right = copy.deepcopy(r)
    r.addnext(right)
    _set_text(r, text[:offset])
    _set_text(right, text[offset:])
    return right


def _wrap_del(r, rid, author, date):
    d = r.makeelement(qn("w:del"), {qn("w:id"): str(rid), qn("w:author"): author, qn("w:date"): date})
    r.addprevious(d)
    d.append(r)
    for t in r.iter(qn("w:t")):
        t.tag = qn("w:delText")
        t.set(XML_SPACE, "preserve")


def _insert_run(after_el, model_run, text, rid, author, date):
    """Tracked insertion of `text`, formatted like model_run, placed before after_el."""
    r = copy.deepcopy(model_run)
    _set_text(r, text)
    ins = r.makeelement(qn("w:ins"), {qn("w:id"): str(rid), qn("w:author"): author, qn("w:date"): date})
    ins.append(r)
    after_el.addprevious(ins)
    return ins


def _apply(p, start, end, next_id, author, date, replacement=""):
    pos, targets = 0, []
    for r, txt in _segments(p):
        a, b = pos, pos + len(txt)
        pos = b
        if b <= start or a >= end:
            continue
        if r is None:
            return None  # span touches a barrier
        targets.append((r, a, b))
    # split boundaries (right edge first so offsets stay valid)
    runs = []
    for r, a, b in targets:
        if b > end:
            _split(r, end - a)
        if a < start:
            r = _split(r, start - a)
        runs.append(r)
    if replacement:
        _insert_run(runs[0], runs[0], replacement, next_id, author, date)
        next_id += 1
    for r in runs:
        _wrap_del(r, next_id, author, date)
        next_id += 1
    return next_id


def paragraph_text(p):
    return "".join(t for _, t in _segments(p))


def track_delete(src, dst, find_spans, author="", date=None):
    doc = Document(src)
    body = doc.element.body
    ids = [int(v) for el in body.iter() for k, v in el.attrib.items() if k == qn("w:id") and v.isdigit()]
    next_id = max(ids, default=0) + 1000
    date = date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    applied, skipped = [], []
    for p in list(body.iter(qn("w:p"))):
        text = paragraph_text(p)
        for span in sorted(find_spans(text), reverse=True):
            start, end = span[0], span[1]
            replacement = span[2] if len(span) > 2 else ""
            snippet = text[start:end]
            if BARRIER in snippet:
                skipped.append(snippet.replace(BARRIER, "·")); continue
            nid = _apply(p, start, end, next_id, author, date, replacement)
            if nid is None:
                skipped.append(snippet)
            else:
                next_id = nid
                applied.append(snippet)
    doc.save(dst)
    return applied, skipped


def views(path):
    """(accepted text, rejected text) per paragraph: accepted keeps insertions and drops deletions,
    rejected keeps deletions and drops insertions."""
    doc = Document(path)
    acc, rej = [], []
    for p in doc.element.body.iter(qn("w:p")):
        a = r = ""
        for el in p.iter():
            if el.tag == qn("w:t"):
                a += el.text or ""
                if not any(anc.tag == qn("w:ins") for anc in el.iterancestors()):
                    r += el.text or ""
            elif el.tag == qn("w:delText"):
                r += el.text or ""
        acc.append(a); rej.append(r)
    return acc, rej


def _selftest():
    import os, tempfile
    d = tempfile.mkdtemp()
    src, dst = os.path.join(d, "a.docx"), os.path.join(d, "b.docx")
    doc = Document()
    p = doc.add_paragraph()
    for chunk, bold in (("ខ្លឹមសារ (Dr", False), ("aft) និង ព័ត៌មាន", True), (" (Draft)។", False)):
        p.add_run(chunk).bold = bold
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "ក្នុងតារាង (Draft)"
    doc.save(src)
    applied, skipped = track_delete(src, dst, lambda t: [m.span() for m in re.finditer(r" ?\(Draft\)", t)], author="T")
    assert len(applied) == 3 and not skipped, (applied, skipped)
    acc, _ = views(dst)
    assert "ខ្លឹមសារ និង ព័ត៌មាន។" in acc and "ក្នុងតារាង" in acc, acc
    orig, _ = views(src)
    _, rej = views(dst)
    assert [x for x in rej if x] == [x for x in orig if x], (rej, orig)
    x = zipfile.ZipFile(dst).read("word/document.xml").decode()
    assert x.count("<w:del ") == 4 and "w:delText" in x  # one span crosses two runs

    # replacement: ប័ណ្ណ -> បណ្ណ keeps formatting and is reversible
    src2, dst2 = os.path.join(d, "c.docx"), os.path.join(d, "e.docx")
    doc2 = Document()
    doc2.add_paragraph().add_run("ប័ណ្ណជំនួយថវិកា និង ប័ណ្ណជំនួយ").bold = True
    doc2.save(src2)
    applied2, skipped2 = track_delete(src2, dst2, lambda t: [(m.start(), m.end(), "បណ្ណ") for m in re.finditer("ប័ណ្ណ", t)], author="T")
    assert len(applied2) == 2 and not skipped2
    acc2, rej2 = views(dst2)
    assert "បណ្ណជំនួយថវិកា និង បណ្ណជំនួយ" in "".join(acc2), acc2
    assert "ប័ណ្ណជំនួយថវិកា និង ប័ណ្ណជំនួយ" in "".join(rej2).replace("បណ្ណ", "", 2), rej2
    y = zipfile.ZipFile(dst2).read("word/document.xml").decode()
    assert y.count("<w:ins ") == 2 and y.count("<w:del ") == 2
    print("selftest ok")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src", nargs="?"); ap.add_argument("dst", nargs="?")
    ap.add_argument("--pattern"); ap.add_argument("--author", default=""); ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return _selftest()
    if not (a.src and a.dst and a.pattern):
        ap.error("IN.docx OUT.docx --pattern are required")
    if a.src == a.dst:
        ap.error("write to a new file")
    rx = re.compile(a.pattern)
    applied, skipped = track_delete(a.src, a.dst, lambda t: [m.span() for m in rx.finditer(t)], a.author)
    print(f"{len(applied)} tracked deletion(s), {len(skipped)} skipped -> {a.dst}")


if __name__ == "__main__":
    sys.exit(main())
